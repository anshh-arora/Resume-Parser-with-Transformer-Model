from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from resume_parser.extractors import UnsupportedFormatError, extract_text


def test_pdf_extraction_pulls_expected_keywords(sample_resume_pdf):
    text = extract_text(sample_resume_pdf)
    # The extracted text should be non-trivial and contain things we know are in the resume.
    assert len(text) > 500
    for needle in ["Ansh Arora", "ansharora.cs@gmail.com", "LangGraph", "Smartworks"]:
        assert needle in text, f"Expected '{needle}' in extracted text"


def test_docx_extraction_round_trip(tmp_path: Path):
    docx_path = tmp_path / "fake.docx"
    doc = Document()
    doc.add_heading("Jane Doe", level=1)
    doc.add_paragraph("Senior Backend Engineer")
    doc.add_paragraph("Skills: Python, Django, PostgreSQL")
    doc.save(str(docx_path))

    text = extract_text(docx_path)
    assert "Jane Doe" in text
    assert "Senior Backend Engineer" in text
    assert "Python" in text


def test_unsupported_extension_raises(tmp_path: Path):
    bad = tmp_path / "resume.txt"
    bad.write_text("plain text resume")
    with pytest.raises(UnsupportedFormatError):
        extract_text(bad)


def test_missing_file_raises(tmp_path: Path):
    with pytest.raises(FileNotFoundError):
        extract_text(tmp_path / "does_not_exist.pdf")
