"""Plain-text extraction from PDF and Word documents."""
from __future__ import annotations

import re
from pathlib import Path

import pdfplumber
from docx import Document

from .config import SUPPORTED_EXTENSIONS
from .logging_setup import get_logger

log = get_logger(__name__)


class UnsupportedFormatError(ValueError):
    pass


def extract_text(file_path: str | Path) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()
    log.info("Extracting text from %s (%s)", path.name, suffix)

    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(
            f"'{suffix}' is not supported. Use one of: {sorted(SUPPORTED_EXTENSIONS)}"
        )
    if not path.exists():
        raise FileNotFoundError(f"Resume file not found: {path}")

    if suffix == ".pdf":
        text = _from_pdf(path)
    elif suffix == ".doc":
        # python-docx cannot read legacy binary .doc (OLE2) files.
        # Try loading via _from_docx; if that fails, give a clear message.
        try:
            text = _from_docx(path)
        except Exception:
            raise UnsupportedFormatError(
                "Legacy .doc (binary) format is not reliably supported. "
                "Please convert to .docx (File → Save As → .docx) or PDF and re-upload."
            )
    else:
        text = _from_docx(path)

    normalized = _normalize(text)
    log.info("Extracted %d characters from %s", len(normalized), path.name)
    return normalized


def _from_pdf(path: Path) -> str:
    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    log.debug("PDF: %d pages", len(pages))
    return "\n".join(pages)


def _from_docx(path: Path) -> str:
    doc = Document(str(path))
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    log.debug("DOCX: %d paragraphs/cells", len(chunks))
    return "\n".join(chunks)


def _normalize(text: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in text.splitlines()]
    return "\n".join(ln for ln in lines if ln)
